from pathlib import Path
import re

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


ROOT = Path('/Users/zhuchong/Documents/code/starrocks-main')
SOURCE = ROOT / 'StarRocks_Rowset_Tenant_TTL_Rewrite_可行性与代码改造方案.docx'
TARGET = ROOT / 'StarRocks_Rowset_Tenant_TTL_Rewrite_可行性与代码改造方案.md'


def iter_blocks(document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, document)
        elif child.tag == qn('w:tbl'):
            yield Table(child, document)


def normalized(text):
    lines = [re.sub(r'[ \t]+', ' ', line).strip() for line in text.splitlines()]
    return '\n'.join(lines).strip()


def escape_table_cell(text):
    text = normalized(text).replace('|', r'\|')
    return '<br>'.join(line for line in text.splitlines() if line)


def paragraph_num_id(paragraph):
    p_pr = paragraph._p.pPr
    if p_pr is None or p_pr.numPr is None or p_pr.numPr.numId is None:
        return None
    return p_pr.numPr.numId.val


def table_to_markdown(table):
    rows = [[escape_table_cell(cell.text) for cell in row.cells] for row in table.rows]
    if not rows:
        return ''

    if len(rows[0]) == 1:
        cell = table.cell(0, 0)
        paragraphs = [p for p in cell.paragraphs if normalized(p.text)]
        if any(p.style and p.style.name == 'Code Block' for p in paragraphs):
            code = '\n'.join(normalized(p.text) for p in paragraphs)
            return f'```text\n{code}\n```'
        if not paragraphs:
            return ''
        title = normalized(paragraphs[0].text)
        body = '\n\n'.join(normalized(p.text) for p in paragraphs[1:])
        lines = [f'> **{title}**']
        if body:
            lines.extend(['>', *[f'> {line}' if line else '>' for line in body.splitlines()]])
        return '\n'.join(lines)

    width = max(len(row) for row in rows)
    rows = [row + [''] * (width - len(row)) for row in rows]
    output = [
        '| ' + ' | '.join(rows[0]) + ' |',
        '| ' + ' | '.join(['---'] * width) + ' |',
    ]
    for row in rows[1:]:
        output.append('| ' + ' | '.join(row) + ' |')
    return '\n'.join(output)


def convert():
    document = Document(SOURCE)
    output = []
    list_counters = {}
    title_seen = 0

    for block in iter_blocks(document):
        if isinstance(block, Table):
            rendered = table_to_markdown(block)
            if rendered:
                output.extend([rendered, ''])
            continue

        text = normalized(block.text)
        if not text:
            continue
        style = block.style.name if block.style else ''

        if text == 'StarRocks' and not output:
            continue
        if style == 'Title':
            title_seen += 1
            output.append(f'# {text}' if title_seen == 1 else f'**{text}**')
        elif style == 'Heading 1':
            output.append(f'# {text}')
        elif style == 'Heading 2':
            output.append(f'## {text}')
        elif style == 'Heading 3':
            output.append(f'### {text}')
        elif style == 'List Bullet':
            output.append(f'- {text}')
        else:
            num_id = paragraph_num_id(block)
            if num_id is not None:
                list_counters[num_id] = list_counters.get(num_id, 0) + 1
                output.append(f'{list_counters[num_id]}. {text}')
            else:
                output.append(text.replace('\n', '<br>'))
        output.append('')

    content = '\n'.join(output)
    content = re.sub(r'\n{3,}', '\n\n', content).strip() + '\n'
    TARGET.write_text(content, encoding='utf-8')
    print(TARGET)


if __name__ == '__main__':
    convert()
